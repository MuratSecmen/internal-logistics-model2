"""
task3_runner.py — MIP vs Solomon karşılaştırması (Task 3).

Deney tasarımı
==============
    Ürün sayısı : 5, 6, 7, 8, 9, 10, 20, 30
    Case        : case1, case2, case3, case4
    Amaç        : route_duration, wait_time
    Araç config : 1 araç, 1 rota
    ─────────────────────────────────────────
    Toplam      : 8 × 4 × 2 = 64 run

Her run için:
    1. Single-objective MIP çözümü (Gurobi)
    2. Solomon çözümü (uygun alpha bias ile)

Çıktı:
    task3_results.xlsx — 64 satır, MIP vs Solomon karşılaştırma sütunları

Solomon alpha seçimi (hoca notu: "alpha_1, alpha_2, alpha_3 objective'e göre"):
    route_duration obj → ("route_duration", alpha_1=0.0001, alpha_2=1.0,  alpha_3=0.01)
    wait_time obj      → ("wait_time",      alpha_1=0.0001, alpha_2=0.01, alpha_3=1.0)
"""

from __future__ import annotations

import copy
import logging
import time
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("internal_logistics.task3")

# ─────────────────────────────────────────────────────────────
# Deney parametreleri
# ─────────────────────────────────────────────────────────────

PRODUCT_COUNTS = [5, 6, 7, 8, 9, 10, 20, 30]
CASES          = ["case1", "case2", "case3", "case4"]
OBJECTIVES     = ["route_duration", "wait_time"]

# Solomon alpha bias (hoca talimatına göre)
_SOLOMON_BIAS = {
    "route_duration": {"alpha_1": 0.0001, "alpha_2": 1.0,  "alpha_3": 0.01},
    "wait_time":      {"alpha_1": 0.0001, "alpha_2": 0.01, "alpha_3": 1.0},
}


# ─────────────────────────────────────────────────────────────
# Sonuç veri yapısı
# ─────────────────────────────────────────────────────────────

@dataclass
class RunResult:
    n_products:  int
    case:        str
    objective:   str

    # MIP
    mip_status:          str   = "not_run"
    mip_primary:         float = float("nan")   # optimize edilen hedef
    mip_secondary:       float = float("nan")   # diğer hedef
    mip_runtime_s:       float = float("nan")
    mip_gap_pct:         float = float("nan")

    # Solomon
    sol_status:          str   = "not_run"
    sol_primary:         float = float("nan")
    sol_secondary:       float = float("nan")
    sol_runtime_s:       float = float("nan")
    sol_bias_used:       str   = ""
    sol_iters:           int   = 0

    # Gap (Solomon - MIP) / MIP, sadece primary hedef için
    gap_pct_primary:     float = float("nan")


# ─────────────────────────────────────────────────────────────
# MIP arayüzü (run_model.py ile bağlantı)
# ─────────────────────────────────────────────────────────────

def _run_mip_single_objective(inst, cfg: dict, objective: str) -> dict:
    """
    Single-objective MIP çözümü çalıştırır.

    run_model.py'deki solve_mip() veya eşdeğer fonksiyonu çağırır.
    objective = "route_duration" → f1 minimize, f2 serbest
    objective = "wait_time"     → f2 minimize, f1 serbest

    Dönüş sözlüğü:
        status         : "optimal" / "feasible" / "infeasible" / "timeout"
        route_duration : float
        total_wait     : float
        runtime_s      : float
        gap_pct        : float (MIP gap yüzdesi)
    """
    # ── Mevcut kod tabanına göre burası güncellenmeli ──
    # run_model.py'den uygun fonksiyonu import edin.
    # Örnek:
    #   from run_model import solve_single_objective_mip
    #   return solve_single_objective_mip(inst, cfg, objective)
    #
    # Geçici stub — entegrasyon sırasında kaldırın:
    raise NotImplementedError(
        "_run_mip_single_objective: run_model.py ile bağlantı kurulmalı.\n"
        "solve_single_objective_mip(inst, cfg, objective) fonksiyonunu import edin."
    )


# ─────────────────────────────────────────────────────────────
# Solomon arayüzü
# ─────────────────────────────────────────────────────────────

def _run_solomon(inst, cfg: dict, objective: str) -> dict:
    """
    Objective'e uygun alpha bias ile Solomon çalıştırır.
    construct_with_multistart kullanır (backtracking + fallback dahil).

    Dönüş:
        status         : "feasible" / "infeasible_unrouted"
        route_duration : float
        total_wait     : float
        runtime_s      : float
        bias_used      : str
        iterations     : int
    """
    from solomon import construct_with_multistart

    sol_cfg = copy.deepcopy(cfg)
    # Objective'e göre primary alpha bias'ı ayarla
    sol_cfg.update(_SOLOMON_BIAS[objective])
    sol_cfg["heuristic_objective"] = objective

    fleet, status, summary = construct_with_multistart(inst, sol_cfg)

    return {
        "status":         status,
        "route_duration": summary.get("route_duration", float("nan")),
        "total_wait":     summary.get("total_wait",     float("nan")),
        "runtime_s":      summary.get("runtime_s",      float("nan")),
        "bias_used":      summary.get("bias_used",      ""),
        "iterations":     summary.get("iterations",     0),
    }


# ─────────────────────────────────────────────────────────────
# Tek run
# ─────────────────────────────────────────────────────────────

def _run_one(inst, cfg: dict, n_products: int,
             case: str, objective: str) -> RunResult:
    res = RunResult(n_products=n_products, case=case, objective=objective)

    # ── MIP ──
    log.info("[Task3] MIP | n=%d  case=%s  obj=%s", n_products, case, objective)
    try:
        mip = _run_mip_single_objective(inst, cfg, objective)
        res.mip_status    = mip["status"]
        res.mip_runtime_s = mip["runtime_s"]
        res.mip_gap_pct   = mip.get("gap_pct", float("nan"))
        if objective == "route_duration":
            res.mip_primary   = mip["route_duration"]
            res.mip_secondary = mip["total_wait"]
        else:
            res.mip_primary   = mip["total_wait"]
            res.mip_secondary = mip["route_duration"]
    except NotImplementedError as e:
        log.warning("MIP stub: %s", e)
        res.mip_status = "stub"
    except Exception as e:
        log.error("MIP hata: %s", e)
        res.mip_status = f"error: {e}"

    # ── Solomon ──
    log.info("[Task3] Solomon | n=%d  case=%s  obj=%s", n_products, case, objective)
    try:
        sol = _run_solomon(inst, cfg, objective)
        res.sol_status    = sol["status"]
        res.sol_runtime_s = sol["runtime_s"]
        res.sol_bias_used = sol["bias_used"]
        res.sol_iters     = sol["iterations"]
        if objective == "route_duration":
            res.sol_primary   = sol["route_duration"]
            res.sol_secondary = sol["total_wait"]
        else:
            res.sol_primary   = sol["total_wait"]
            res.sol_secondary = sol["route_duration"]
    except Exception as e:
        log.error("Solomon hata: %s", e)
        res.sol_status = f"error: {e}"

    # ── Gap hesabı ──
    import math
    if (not math.isnan(res.mip_primary) and not math.isnan(res.sol_primary)
            and res.mip_primary != 0):
        res.gap_pct_primary = (
            (res.sol_primary - res.mip_primary) / abs(res.mip_primary) * 100
        )

    log.info(
        "[Task3] DONE | n=%d case=%s obj=%s | MIP=%.2f Sol=%.2f gap=%.1f%%",
        n_products, case, objective,
        res.mip_primary, res.sol_primary, res.gap_pct_primary,
    )
    return res


# ─────────────────────────────────────────────────────────────
# Ana runner
# ─────────────────────────────────────────────────────────────

def run_task3(cfg: dict,
              product_counts=None,
              cases=None,
              objectives=None,
              output_path: str = "task3_results.xlsx") -> list[RunResult]:
    """
    64 run'ı (veya alt kümeyi) çalıştırır ve Excel dosyası üretir.

    Parametreler
    ------------
    cfg            : temel konfigürasyon sözlüğü (run_model.py'deki gibi)
    product_counts : test edilecek ürün sayıları; None → tüm 8 değer
    cases          : test edilecek case'ler; None → 4 case
    objectives     : hedef fonksiyonları; None → her ikisi
    output_path    : çıktı Excel dosyasının yolu

    Dönüş
    -----
    list[RunResult] — tüm run sonuçları
    """
    # ── Mevcut kod tabanına göre load_instance import edilmeli ──
    # from run_model import load_instance
    try:
        from run_model import load_instance
    except ImportError:
        log.error("run_model.load_instance bulunamadı; stub kullanılıyor.")
        load_instance = None

    p_counts  = product_counts or PRODUCT_COUNTS
    cs        = cases          or CASES
    objs      = objectives     or OBJECTIVES

    total   = len(p_counts) * len(cs) * len(objs)
    results = []
    run_no  = 0

    for n in p_counts:
        for case in cs:
            for obj in objs:
                run_no += 1
                log.info("── Run %d/%d: n=%d  case=%s  obj=%s ──",
                         run_no, total, n, case, obj)

                # Instance yükle
                if load_instance:
                    run_cfg = copy.deepcopy(cfg)
                    run_cfg["n_products"] = n
                    try:
                        inst = load_instance(case, n, run_cfg)
                    except Exception as e:
                        log.error("Instance yüklenemedi: %s", e)
                        results.append(RunResult(n, case, obj,
                                                 mip_status=f"load_error:{e}",
                                                 sol_status=f"load_error:{e}"))
                        continue
                else:
                    inst = None
                    run_cfg = cfg

                r = _run_one(inst, run_cfg, n, case, obj)
                results.append(r)

    _write_excel(results, output_path)
    log.info("Task3 tamamlandı. Sonuçlar: %s", output_path)
    return results


# ─────────────────────────────────────────────────────────────
# Excel yazıcı
# ─────────────────────────────────────────────────────────────

_COL_HEADERS = [
    ("n_products",       "Ürün\nSayısı"),
    ("case",             "Case"),
    ("objective",        "Amaç\nFonksiyonu"),
    ("mip_status",       "MIP\nDurum"),
    ("mip_primary",      "MIP\nPrimary Obj"),
    ("mip_secondary",    "MIP\nSecondary Obj"),
    ("mip_runtime_s",    "MIP\nSüre (s)"),
    ("mip_gap_pct",      "MIP\nGap (%)"),
    ("sol_status",       "Solomon\nDurum"),
    ("sol_primary",      "Solomon\nPrimary Obj"),
    ("sol_secondary",    "Solomon\nSecondary Obj"),
    ("sol_runtime_s",    "Solomon\nSüre (s)"),
    ("sol_bias_used",    "Solomon\nBias"),
    ("sol_iters",        "Solomon\nİter"),
    ("gap_pct_primary",  "Gap\n(Sol-MIP)/MIP %"),
]

_HEADER_FILL  = PatternFill("solid", fgColor="2E4057")
_SUBHEAD_FILL = PatternFill("solid", fgColor="4A6FA5")
_GOOD_FILL    = PatternFill("solid", fgColor="C8E6C9")
_BAD_FILL     = PatternFill("solid", fgColor="FFCDD2")
_ALT_FILL     = PatternFill("solid", fgColor="F5F5F5")

_FLOAT_FMT = "0.00"
_PCT_FMT   = "0.0%"


def _write_excel(results: list[RunResult], path: str) -> None:
    """64 satırlık özet Excel dosyası yazar."""
    import math

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIP vs Solomon"

    # ── Başlık satırı ──
    ws.merge_cells("A1:O1")
    title_cell = ws["A1"]
    title_cell.value    = "Task 3 — MIP vs Solomon Karşılaştırması"
    title_cell.font     = Font(bold=True, color="FFFFFF", size=13)
    title_cell.fill     = _HEADER_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # ── Sütun başlıkları (satır 2) ──
    for col_idx, (_, header) in enumerate(_COL_HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = _SUBHEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[2].height = 30

    # ── Veri satırları ──
    for row_idx, r in enumerate(results, start=3):
        alt = (row_idx % 2 == 0)
        values = [
            r.n_products, r.case, r.objective,
            r.mip_status,
            r.mip_primary,   r.mip_secondary,
            r.mip_runtime_s, r.mip_gap_pct,
            r.sol_status,
            r.sol_primary,   r.sol_secondary,
            r.sol_runtime_s, r.sol_bias_used,
            r.sol_iters,
            r.gap_pct_primary,
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # Sayısal biçimlendirme
            attr = _COL_HEADERS[col_idx - 1][0]
            if attr in ("mip_primary", "mip_secondary",
                        "sol_primary", "sol_secondary"):
                cell.number_format = "0.000"
            elif attr in ("mip_runtime_s", "sol_runtime_s", "mip_gap_pct"):
                cell.number_format = "0.00"
            elif attr == "gap_pct_primary":
                cell.number_format = "0.0"

            # Gap renklendirmesi
            if attr == "gap_pct_primary" and isinstance(val, float) and not math.isnan(val):
                cell.fill = _GOOD_FILL if val <= 5.0 else _BAD_FILL
            elif alt and attr not in ("gap_pct_primary",):
                cell.fill = _ALT_FILL

    # ── Sütun genişlikleri ──
    widths = [10, 8, 16, 12, 14, 14, 11, 10, 12, 14, 14, 11, 14, 8, 16]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # ── Özet istatistik sayfası ──
    ws2 = wb.create_sheet("Özet")
    _write_summary_sheet(ws2, results)

    wb.save(path)
    log.info("Excel kaydedildi: %s", path)


def _write_summary_sheet(ws, results: list[RunResult]) -> None:
    """Objective × case bazında ortalama gap özeti."""
    import math

    ws["A1"] = "Objective"
    ws["B1"] = "Case"
    ws["C1"] = "Ort. Gap % (Primary)"
    ws["D1"] = "Feasible Run Sayısı"
    ws["E1"] = "Ortalama MIP Süre (s)"
    ws["F1"] = "Ortalama Solomon Süre (s)"

    for cell in ws["A1:F1"][0]:
        cell.font = Font(bold=True)
        cell.fill = _SUBHEAD_FILL
        cell.font = Font(bold=True, color="FFFFFF")

    row = 2
    for obj in OBJECTIVES:
        for case in CASES:
            subset = [r for r in results
                      if r.objective == obj and r.case == case]
            gaps = [r.gap_pct_primary for r in subset
                    if not math.isnan(r.gap_pct_primary)]
            mip_times = [r.mip_runtime_s for r in subset
                         if not math.isnan(r.mip_runtime_s)]
            sol_times = [r.sol_runtime_s for r in subset
                         if not math.isnan(r.sol_runtime_s)]
            feasible = sum(1 for r in subset
                           if r.mip_status in ("optimal", "feasible")
                           and r.sol_status == "feasible")

            ws.cell(row=row, column=1, value=obj)
            ws.cell(row=row, column=2, value=case)
            ws.cell(row=row, column=3,
                    value=(sum(gaps) / len(gaps) if gaps else float("nan")))
            ws.cell(row=row, column=4, value=feasible)
            ws.cell(row=row, column=5,
                    value=(sum(mip_times) / len(mip_times)
                           if mip_times else float("nan")))
            ws.cell(row=row, column=6,
                    value=(sum(sol_times) / len(sol_times)
                           if sol_times else float("nan")))
            row += 1

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 8
    for col in ["C", "D", "E", "F"]:
        ws.column_dimensions[col].width = 22


# ─────────────────────────────────────────────────────────────
# CLI giriş noktası
# ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import json, sys

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )

    cfg_path = sys.argv[1] if len(sys.argv) > 1 else "config.json"
    with open(cfg_path) as f:
        cfg = json.load(f)

    results = run_task3(cfg)
    print(f"\nTask 3 tamamlandı: {len(results)} run")
    feasible_mip = sum(1 for r in results if r.mip_status in ("optimal","feasible"))
    feasible_sol = sum(1 for r in results if r.sol_status == "feasible")
    print(f"  Feasible MIP    : {feasible_mip}/{len(results)}")
    print(f"  Feasible Solomon: {feasible_sol}/{len(results)}")
