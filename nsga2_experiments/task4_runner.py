"""
task4_runner.py — MIP (Augmented ε-constraint) vs NSGA-2 karşılaştırması (Task 4).

Deney tasarımı
==============
    Ürün sayısı : 5, 6, 7, 8, 9, 10
    Case        : case1, case2, case3, case4
    Araç config : (1 araç, 1 rota), (3 araç, 3 rota)
    ─────────────────────────────────────────────────
    Toplam      : 6 × 4 × 2 = 48 run

MIP tarafı — Augmented ε-constraint sweep
    Ana amaç   : min route_duration (f1)
    ε-kısıt    : total_wait ≤ ε (f2)
    Augmented  : min f1 + 10⁻³ × f2   s.t. f2 ≤ ε
    Sweep      : ε_0 = ∞ (kısıtsız) → ε_{t+1} = wait_t − 5 → infeasible'a kadar

NSGA-2 tarafı
    Her iki objective birlikte (zaten bi-objective yapı)
    Solomon-tabanlı decoder ile

Karşılaştırma metrikleri (frontier_metrics.py):
    HV, GD, IGD, SP, EI

Çıktılar:
    task4_results.xlsx     — 48 satır özet + metrik sütunları
    plots/                 — her run için MIP vs NSGA-2 Pareto plot (PNG)
"""

from __future__ import annotations

import copy
import json
import logging
import math
import os
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from frontier_metrics import compute_all

log = logging.getLogger("internal_logistics.task4")

# ─────────────────────────────────────────────────────────────
# Deney parametreleri
# ─────────────────────────────────────────────────────────────

PRODUCT_COUNTS = [5, 6, 7, 8, 9, 10]
CASES          = ["case1", "case2", "case3", "case4"]
VEHICLE_CONFIGS = [
    {"label": "1V_1R", "n_vehicles": 1, "max_routes": 1},
    {"label": "3V_3R", "n_vehicles": 3, "max_routes": 3},
]
EPSILON_STEP   = 5      # ε her adımda 5 azalır
EPSILON_MIN    = 0      # ε bu değerin altına inmesin


# ─────────────────────────────────────────────────────────────
# Sonuç veri yapısı
# ─────────────────────────────────────────────────────────────

@dataclass
class Task4Result:
    n_products:    int
    case:          str
    vehicle_label: str

    # MIP Pareto frontu — [(route_duration, total_wait), ...]
    mip_front:   list = field(default_factory=list)
    mip_n_eps:   int  = 0       # kaç ε noktası çözüldü
    mip_runtime_total_s: float = float("nan")

    # NSGA-2 Pareto frontu
    nsga2_front:          list  = field(default_factory=list)
    nsga2_n_generations:  int   = 0
    nsga2_runtime_s:      float = float("nan")

    # Metrikler (NSGA-2 frontu, referans = MIP frontu)
    hv:    float = float("nan")
    gd:    float = float("nan")
    igd:   float = float("nan")
    sp:    float = float("nan")
    ei:    float = float("nan")

    # Hata durumları
    mip_error:   str = ""
    nsga2_error: str = ""


# ─────────────────────────────────────────────────────────────
# MIP Augmented ε-constraint sweep
# ─────────────────────────────────────────────────────────────

def _run_augmented_epsilon_mip(inst, cfg: dict,
                                step: int = EPSILON_STEP) -> dict:
    """
    Augmented ε-constraint sweep çalıştırır.

    Algoritma:
        1. ε = ∞ → kısıtsız min route_duration (augmented: +10⁻³×wait)
           Çözüm: (dur_0, wait_0)
        2. ε = wait_0 − step → tekrar çöz → (dur_1, wait_1)
        3. ... → infeasible veya wait ≤ EPSILON_MIN'e kadar

    Her adımda Gurobi'ye augmented objective verilir:
        min route_duration + 1e-3 × total_wait
        s.t. total_wait ≤ ε
             + tüm MIP kısıtları (3)-(35)

    Dönüş:
        {
          "front"     : [(dur, wait), ...],  # Pareto noktaları
          "n_points"  : int,
          "runtime_s" : float,               # toplam MIP süresi
          "status"    : "ok" / "error"
        }
    """
    # ── run_model.py bağlantısı ──
    # from run_model import solve_augmented_epsilon_mip
    # Bu fonksiyon şu signature'ı bekleniyor:
    #   solve_augmented_epsilon_mip(inst, cfg, epsilon_wait) -> dict
    #   dönen dict: {status, route_duration, total_wait, runtime_s, gap_pct}
    #
    # Stub — entegrasyon sırasında kaldırın:
    try:
        from run_model import solve_augmented_epsilon_mip
    except ImportError:
        raise NotImplementedError(
            "run_model.solve_augmented_epsilon_mip fonksiyonu bulunamadı.\n"
            "run_model.py'ye şu fonksiyonu ekleyin:\n"
            "  def solve_augmented_epsilon_mip(inst, cfg, epsilon_wait: float) -> dict\n"
            "  Dönüş: {status, route_duration, total_wait, runtime_s, gap_pct}"
        )

    front     = []
    t_total   = 0.0
    epsilon   = float("inf")   # ilk çözüm kısıtsız

    while True:
        run_cfg = copy.deepcopy(cfg)
        run_cfg["epsilon_wait"] = epsilon

        t0  = time.perf_counter()
        sol = solve_augmented_epsilon_mip(inst, run_cfg, epsilon)
        t0  = time.perf_counter() - t0
        t_total += t0

        if sol["status"] in ("infeasible", "error"):
            log.info("  ε=%.1f → %s; sweep bitti", epsilon, sol["status"])
            break

        dur  = sol["route_duration"]
        wait = sol["total_wait"]
        log.info("  ε=%.1f → dur=%.2f  wait=%.2f  gap=%.1f%%",
                 epsilon, dur, wait, sol.get("gap_pct", float("nan")))

        # Yineleme varsa ekleme (epsilon hızla yakınsıyorsa)
        if not any(abs(p[0] - dur) < 1e-4 and abs(p[1] - wait) < 1e-4
                   for p in front):
            front.append((dur, wait))

        # Sonraki epsilon
        epsilon = wait - step
        if epsilon < EPSILON_MIN:
            log.info("  ε=%.1f < min=%d; sweep bitti", epsilon, EPSILON_MIN)
            break

    return {
        "front":     front,
        "n_points":  len(front),
        "runtime_s": t_total,
        "status":    "ok" if front else "no_solution",
    }


# ─────────────────────────────────────────────────────────────
# NSGA-2 çalıştırıcı
# ─────────────────────────────────────────────────────────────

def _run_nsga2(inst, cfg: dict) -> dict:
    """
    NSGA-2'yi çalıştırır; Pareto frontunu döner.

    Dönüş:
        {
          "front"        : [(route_duration, total_wait), ...],
          "n_generations": int,
          "runtime_s"    : float,
          "status"       : "ok" / "infeasible" / "error"
        }
    """
    from nsga2 import nsga2

    t0 = time.perf_counter()
    pareto, history = nsga2(inst, cfg)
    runtime = time.perf_counter() - t0

    if not pareto:
        return {
            "front": [], "n_generations": len(history),
            "runtime_s": runtime, "status": "infeasible",
        }

    front = [(ind.f1, ind.f2) for ind in pareto]
    return {
        "front":         front,
        "n_generations": len(history),
        "runtime_s":     runtime,
        "status":        "ok",
    }


# ─────────────────────────────────────────────────────────────
# Tek run
# ─────────────────────────────────────────────────────────────

def _run_one(inst, cfg: dict, n: int, case: str,
             vcfg: dict, plot_dir: Path) -> Task4Result:
    label = vcfg["label"]
    res   = Task4Result(n_products=n, case=case, vehicle_label=label)

    # ── MIP sweep ──
    log.info("[Task4] MIP ε-sweep | n=%d  case=%s  %s", n, case, label)
    try:
        mip = _run_augmented_epsilon_mip(inst, cfg)
        res.mip_front          = mip["front"]
        res.mip_n_eps          = mip["n_points"]
        res.mip_runtime_total_s = mip["runtime_s"]
    except NotImplementedError as e:
        log.warning("MIP stub: %s", e)
        res.mip_error = "stub"
    except Exception as e:
        log.error("MIP hata: %s", e)
        res.mip_error = str(e)

    # ── NSGA-2 ──
    log.info("[Task4] NSGA-2 | n=%d  case=%s  %s", n, case, label)
    try:
        ns = _run_nsga2(inst, cfg)
        res.nsga2_front         = ns["front"]
        res.nsga2_n_generations = ns["n_generations"]
        res.nsga2_runtime_s     = ns["runtime_s"]
    except Exception as e:
        log.error("NSGA-2 hata: %s", e)
        res.nsga2_error = str(e)

    # ── Metrikler ──
    if res.mip_front and res.nsga2_front:
        all_pts   = res.mip_front + res.nsga2_front
        ref_point = (
            max(p[0] for p in all_pts) * 1.10 + 1,
            max(p[1] for p in all_pts) * 1.10 + 1,
        )
        metrics = compute_all(
            approx    = res.nsga2_front,
            ref_set   = res.mip_front,
            ref_point = ref_point,
            normalize = True,
        )
        res.hv  = metrics["HV"]
        res.gd  = metrics["GD"]
        res.igd = metrics["IGD"]
        res.sp  = metrics["SP"]
        res.ei  = metrics["EI"]
        log.info(
            "  Metrikler | HV=%.4f  GD=%.4f  IGD=%.4f  SP=%.4f  EI=%.4f",
            res.hv, res.gd, res.igd, res.sp, res.ei,
        )

    # ── Plot ──
    plot_path = plot_dir / f"pareto_{n}p_{case}_{label}.png"
    _plot_pareto(res, plot_path)

    return res


# ─────────────────────────────────────────────────────────────
# Pareto plot
# ─────────────────────────────────────────────────────────────

def _plot_pareto(res: Task4Result, path: Path) -> None:
    """
    MIP ve NSGA-2 Pareto frontlarını aynı grafikte gösterir.
    matplotlib gerektirir.
    """
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import matplotlib.lines as mlines
    except ImportError:
        log.warning("matplotlib bulunamadı; plot atlandı.")
        return

    fig, ax = plt.subplots(figsize=(7, 5))

    # MIP front
    if res.mip_front:
        mip_sorted = sorted(res.mip_front, key=lambda p: p[0])
        mip_x = [p[0] for p in mip_sorted]
        mip_y = [p[1] for p in mip_sorted]
        ax.plot(mip_x, mip_y, "o-", color="#2E4057", linewidth=1.5,
                markersize=7, label="MIP (augmented ε)", zorder=3)

    # NSGA-2 front
    if res.nsga2_front:
        ns_sorted = sorted(res.nsga2_front, key=lambda p: p[0])
        ns_x = [p[0] for p in ns_sorted]
        ns_y = [p[1] for p in ns_sorted]
        ax.plot(ns_x, ns_y, "s--", color="#D85A30", linewidth=1.5,
                markersize=7, label="NSGA-2 (Solomon decoder)", zorder=3)

    # Metrik annotasyonu
    metric_txt = (
        f"HV={res.hv:.4f}  GD={res.gd:.4f}\n"
        f"IGD={res.igd:.4f}  SP={res.sp:.4f}  EI={res.ei:.4f}"
        if not math.isnan(res.hv) else "Metrik hesaplanamadı"
    )
    ax.text(0.02, 0.97, metric_txt, transform=ax.transAxes,
            fontsize=8, verticalalignment="top",
            bbox=dict(boxstyle="round", facecolor="wheat", alpha=0.4))

    ax.set_xlabel("Route Duration (f₁)", fontsize=11)
    ax.set_ylabel("Total Wait Time (f₂)", fontsize=11)
    ax.set_title(
        f"Pareto Front — n={res.n_products}  {res.case}  {res.vehicle_label}",
        fontsize=12,
    )
    ax.legend(fontsize=9)
    ax.grid(True, alpha=0.3, linewidth=0.5)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    log.info("  Plot: %s", path)


# ─────────────────────────────────────────────────────────────
# Ana runner
# ─────────────────────────────────────────────────────────────

def run_task4(cfg: dict,
              product_counts=None,
              cases=None,
              vehicle_configs=None,
              output_path: str = "task4_results.xlsx",
              plot_dir: str   = "plots") -> list[Task4Result]:
    """
    48 run çalıştırır; Excel + PNG plotlar üretir.

    Parametreler
    ------------
    cfg             : temel konfigürasyon
    product_counts  : None → [5..10]
    cases           : None → 4 case
    vehicle_configs : None → [(1V,1R), (3V,3R)]
    output_path     : Excel çıktısı
    plot_dir        : PNG plotların dizini
    """
    try:
        from run_model import load_instance
    except ImportError:
        log.error("run_model.load_instance bulunamadı.")
        load_instance = None

    p_counts = product_counts or PRODUCT_COUNTS
    cs       = cases          or CASES
    vcs      = vehicle_configs or VEHICLE_CONFIGS

    plot_path = Path(plot_dir)
    plot_path.mkdir(parents=True, exist_ok=True)

    total   = len(p_counts) * len(cs) * len(vcs)
    results = []
    run_no  = 0

    for n in p_counts:
        for case in cs:
            for vcfg in vcs:
                run_no += 1
                log.info("══ Run %d/%d: n=%d  case=%s  %s ══",
                         run_no, total, n, case, vcfg["label"])

                run_cfg = copy.deepcopy(cfg)
                run_cfg["n_products"] = n
                # Araç config'ini aktif cfg'ye yaz
                run_cfg.setdefault("nsga2", {})
                # (run_model.py'nin beklediği araç parametrelerine göre uyarlayın)
                run_cfg["_vehicle_override"] = vcfg

                if load_instance:
                    try:
                        inst = load_instance(case, n, run_cfg)
                    except Exception as e:
                        log.error("Instance yüklenemedi: %s", e)
                        results.append(Task4Result(n, case, vcfg["label"],
                                                    mip_error=str(e),
                                                    nsga2_error=str(e)))
                        continue
                else:
                    inst = None

                r = _run_one(inst, run_cfg, n, case, vcfg, plot_path)
                results.append(r)

    _write_excel_task4(results, output_path)
    log.info("Task4 tamamlandı. Sonuçlar: %s | Plotlar: %s/",
             output_path, plot_dir)
    return results


# ─────────────────────────────────────────────────────────────
# Excel yazıcı — Task 4
# ─────────────────────────────────────────────────────────────

_T4_HEADERS = [
    ("n_products",           "Ürün\nSayısı"),
    ("case",                 "Case"),
    ("vehicle_label",        "Araç\nConfig"),
    ("mip_n_eps",            "MIP\nε Nokta"),
    ("mip_runtime_total_s",  "MIP\nToplam Süre (s)"),
    ("nsga2_n_generations",  "NSGA-2\nNesil"),
    ("nsga2_runtime_s",      "NSGA-2\nSüre (s)"),
    ("hv",                   "HV\n(↑ iyi)"),
    ("gd",                   "GD\n(↓ iyi)"),
    ("igd",                  "IGD\n(↓ iyi)"),
    ("sp",                   "SP\n(↓ düzgün)"),
    ("ei",                   "ε-ind\n(↓ iyi)"),
    ("mip_error",            "MIP\nHata"),
    ("nsga2_error",          "NSGA-2\nHata"),
]

_H_FILL  = PatternFill("solid", fgColor="2E4057")
_SH_FILL = PatternFill("solid", fgColor="4A6FA5")
_ALT     = PatternFill("solid", fgColor="F5F5F5")
_GOOD    = PatternFill("solid", fgColor="C8E6C9")
_MED     = PatternFill("solid", fgColor="FFF9C4")
_BAD     = PatternFill("solid", fgColor="FFCDD2")


def _write_excel_task4(results: list[Task4Result], path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MIP vs NSGA-2"

    # Başlık
    n_cols = len(_T4_HEADERS)
    ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
    c = ws["A1"]
    c.value     = "Task 4 — MIP (Augmented ε) vs NSGA-2 Karşılaştırması"
    c.font      = Font(bold=True, color="FFFFFF", size=13)
    c.fill      = _H_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Sütun başlıkları
    for ci, (_, hdr) in enumerate(_T4_HEADERS, 1):
        cell = ws.cell(row=2, column=ci, value=hdr)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = _SH_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[2].height = 30

    # Veri satırları
    for ri, r in enumerate(results, 3):
        alt = ri % 2 == 0
        row_data = [
            r.n_products, r.case, r.vehicle_label,
            r.mip_n_eps,  r.mip_runtime_total_s,
            r.nsga2_n_generations, r.nsga2_runtime_s,
            r.hv,  r.gd, r.igd, r.sp, r.ei,
            r.mip_error, r.nsga2_error,
        ]
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            attr = _T4_HEADERS[ci - 1][0]

            if attr in ("mip_runtime_total_s", "nsga2_runtime_s"):
                cell.number_format = "0.00"
            elif attr in ("hv", "gd", "igd", "sp", "ei"):
                cell.number_format = "0.0000"
                # Renklendirme: IGD ve GD için küçük iyi
                if not (isinstance(val, float) and math.isnan(val)):
                    if attr == "hv":
                        cell.fill = _GOOD  # yüksek HV her zaman iyi
                    elif attr in ("gd", "igd", "sp"):
                        cell.fill = _GOOD if val < 0.05 else (
                            _MED if val < 0.15 else _BAD)
                    elif attr == "ei":
                        cell.fill = _GOOD if val < 0 else (
                            _MED if val < 0.05 else _BAD)
            elif alt and attr not in ("hv","gd","igd","sp","ei"):
                cell.fill = _ALT

    # Sütun genişlikleri
    widths = [9, 8, 10, 9, 17, 10, 14, 10, 10, 10, 10, 10, 14, 14]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # ── Metrik özet sayfası ──
    ws2 = wb.create_sheet("Metrik Özeti")
    _write_metric_summary(ws2, results)

    wb.save(path)
    log.info("Excel kaydedildi: %s", path)


def _write_metric_summary(ws, results: list[Task4Result]) -> None:
    """Araç config × case bazında ortalama metrikler."""
    headers = ["Araç Config", "Case",
               "Ort. HV", "Ort. GD", "Ort. IGD", "Ort. SP", "Ort. EI",
               "Feasible Run"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _SH_FILL
        ws.column_dimensions[get_column_letter(ci)].width = 14

    row = 2
    for vcfg in VEHICLE_CONFIGS:
        for case in CASES:
            sub = [r for r in results
                   if r.vehicle_label == vcfg["label"] and r.case == case]
            def avg(attr):
                vals = [getattr(r, attr) for r in sub
                        if not math.isnan(getattr(r, attr))]
                return sum(vals) / len(vals) if vals else float("nan")

            feasible = sum(1 for r in sub if r.mip_front and r.nsga2_front)
            ws.cell(row=row, column=1, value=vcfg["label"])
            ws.cell(row=row, column=2, value=case)
            ws.cell(row=row, column=3, value=avg("hv")).number_format  = "0.0000"
            ws.cell(row=row, column=4, value=avg("gd")).number_format  = "0.0000"
            ws.cell(row=row, column=5, value=avg("igd")).number_format = "0.0000"
            ws.cell(row=row, column=6, value=avg("sp")).number_format  = "0.0000"
            ws.cell(row=row, column=7, value=avg("ei")).number_format  = "0.0000"
            ws.cell(row=row, column=8, value=feasible)
            row += 1


# ─────────────────────────────────────────────────────────────
# CLI giriş noktası
# ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import sys

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )

    cfg_path = sys.argv[1] if len(sys.argv) > 1 else "config.json"
    with open(cfg_path) as f:
        cfg = json.load(f)

    results = run_task4(cfg)
    print(f"\nTask 4 tamamlandı: {len(results)} run")
    feasible = sum(1 for r in results if r.mip_front and r.nsga2_front)
    print(f"  Her iki taraf da feasible: {feasible}/{len(results)}")
