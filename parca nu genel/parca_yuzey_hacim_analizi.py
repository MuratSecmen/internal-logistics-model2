"""
parca_yuzey_hacim_analizi.py
============================
2025 ve 2026 parça taşımacılığı verilerinden:
  - Her MALZEMENO için en yüksek 2 yüzey alanı (WIDTH × LENGTH, cm²)
  - Hacim (GAGE × WIDTH × LENGTH, cm³)

Kullanım:
  pip install pandas openpyxl
  python parca_yuzey_hacim_analizi.py

Dosya yollarını aşağıdaki DEĞİŞKENLERDEN güncelleyebilirsiniz.
"""

import os
import pandas as pd
import openpyxl
import warnings
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

# ════════════════════════════════════════════════════════════
# ▶ KULLANICI AYARLARI — sadece bu bölümü düzenleyin
# ════════════════════════════════════════════════════════════

# Varsayılan olarak bu script ile aynı klasördeki dosyalar kullanılır.
# Farklı bir konum kullanmak isterseniz aşağıdaki üç satırı güncelleyin.
_here = os.path.dirname(os.path.abspath(__file__))
F25 = os.path.join(_here, "2025_parça detayları.xlsx")
F26 = os.path.join(_here, "2026_taşınan parça nu.s.xlsx")

# Çıktı dosyasının kaydedileceği yer
OUT = os.path.join(_here, "parca_yuzey_hacim_analizi.xlsx")

# ── Birim ayarları ──────────────────────────────────────────
# Amerikalı projeler (inç kullananlar) — GAGE=0.032 kanıtıyla tespit edildi
# Bildiğiniz başka Amerikalı proje kodlarını buraya ekleyin:
INCH_PROJECTS = {'GMH', 'LGM'}

# Eğer WIDTH veya LENGTH bu değerden büyükse → mm kabul edilir, ÷10 ile cm'e çevrilir
# (Örn: WIDTH=1220 → 1220mm = 122cm anlamına gelir; 1220cm = 12.2m imkânsız)
MM_THRESHOLD = 500

INCH_TO_CM = 2.54   # 1 inç = 2.54 cm
MM_TO_CM   = 0.10   # 1 mm  = 0.10 cm

# ════════════════════════════════════════════════════════════
# Buradan aşağısını değiştirmenize gerek yok
# ════════════════════════════════════════════════════════════

print("=" * 60)
print("Parça Yüzey Alanı & Hacim Analizi")
print("=" * 60)

# ─── 1. Veri yükleme & temizlik ──────────────────────────────
print("\n[1/5] Dosyalar okunuyor...")
df25 = pd.read_excel(F25, sheet_name='son')
df26 = pd.read_excel(F26, sheet_name='PART NU_İLK 5 AY')

def hazirla(df, kaynak):
    df = df.copy()
    df['PROJE']     = df['PROJE'].astype(str).str.strip()
    df['MALZEMENO'] = df['MALZEMENO'].astype(str).str.strip()
    df['KAYNAK']    = kaynak
    for c in ['GAGE', 'WIDTH', 'LENGTH']:
        df[c] = pd.to_numeric(df[c], errors='coerce')
    # Test numunelerini filtrele (BONDTEST vb. — GAGE/WIDTH/LENGTH çok küçük)
    df = df[(df['GAGE'] >= 0.01) & (df['WIDTH'] >= 0.1) & (df['LENGTH'] >= 0.1)]
    return df.dropna(subset=['GAGE', 'WIDTH', 'LENGTH', 'MALZEMENO'])

df25 = hazirla(df25, '2025')
df26 = hazirla(df26, '2026')
print(f"  2025 → {len(df25):,} satır")
print(f"  2026 → {len(df26):,} satır")

# ─── 2. Birim tespiti & dönüşüm ──────────────────────────────
print("\n[2/5] Birim tespiti yapılıyor...")

def birim_donustur(row):
    """
    3 kural:
      1. PROJE inç listesindeyse   → tüm boyutlar × 2.54
      2. WIDTH veya LENGTH > 500   → mm kabul et, ÷ 10
      3. Diğerleri                 → zaten cm, dokunma
    """
    p = row['PROJE']
    g, w, l = row['GAGE'], row['WIDTH'], row['LENGTH']
    if p in INCH_PROJECTS:
        return g * INCH_TO_CM, w * INCH_TO_CM, l * INCH_TO_CM, 'inç→cm'
    elif w > MM_THRESHOLD or l > MM_THRESHOLD:
        return g * MM_TO_CM, w * MM_TO_CM, l * MM_TO_CM, 'mm→cm'
    else:
        return g, w, l, 'cm'

for df in [df25, df26]:
    sonuc = df.apply(birim_donustur, axis=1, result_type='expand')
    sonuc.columns = ['GAGE_cm', 'WIDTH_cm', 'LENGTH_cm', 'BIRIM']
    df[['GAGE_cm', 'WIDTH_cm', 'LENGTH_cm', 'BIRIM']] = sonuc
    df['YUZEY_cm2'] = (df['WIDTH_cm'] * df['LENGTH_cm']).round(2)
    df['HACIM_cm3'] = (df['GAGE_cm']  * df['WIDTH_cm'] * df['LENGTH_cm']).round(2)

# Birim özeti
for name, df in [('2025', df25), ('2026', df26)]:
    ozet = df.groupby('BIRIM').size().reset_index(name='Satır')
    print(f"  {name}: " + "  |  ".join(
        f"{r['BIRIM']}:{r['Satır']:,}" for _, r in ozet.iterrows()))

# ─── 3. Her MALZEMENO için top-2 yüzey alanı ─────────────────
print("\n[3/5] Her MALZEMENO için en yüksek 2 yüzey alanı hesaplanıyor...")

def top2_per_malzeme(df):
    return (
        df.sort_values('YUZEY_cm2', ascending=False)
          .groupby('MALZEMENO', group_keys=False)
          .head(2)
          .assign(SIRA=lambda x: x.groupby('MALZEMENO').cumcount() + 1)
          .sort_values(['MALZEMENO', 'SIRA'])
    )

res25 = top2_per_malzeme(df25)
res26 = top2_per_malzeme(df26)
print(f"  2025 → {res25['MALZEMENO'].nunique()} unique MALZEMENO, {len(res25)} satır")
print(f"  2026 → {res26['MALZEMENO'].nunique()} unique MALZEMENO, {len(res26)} satır")

# ─── 4. Özet (MALZEMENO başına max değerler) ──────────────────
print("\n[4/5] Özet tablo oluşturuluyor...")

def ozet_tablo(df):
    return (
        df.groupby('MALZEMENO').agg(
            PROJE          = ('PROJE',    lambda x: ', '.join(sorted(set(x)))),
            BIRIM          = ('BIRIM',    'first'),
            GAGE_MAX_cm    = ('GAGE_cm',  'max'),
            WIDTH_MAX_cm   = ('WIDTH_cm', 'max'),
            LENGTH_MAX_cm  = ('LENGTH_cm','max'),
            YUZEY_MAX_cm2  = ('YUZEY_cm2','max'),
            HACIM_MAX_cm3  = ('HACIM_cm3','max'),
            SATIR_N        = ('GAGE',     'count'),
        )
        .reset_index()
        .sort_values('YUZEY_MAX_cm2', ascending=False)
    )

oz25 = ozet_tablo(df25)
oz26 = ozet_tablo(df26)

# Konsola ilk 5'i göster
print("\n  ── 2025 Top-5 ──")
print(oz25[['MALZEMENO','BIRIM','WIDTH_MAX_cm','LENGTH_MAX_cm','YUZEY_MAX_cm2','HACIM_MAX_cm3']].head(5).to_string(index=False))
print("\n  ── 2026 Top-5 ──")
print(oz26[['MALZEMENO','BIRIM','WIDTH_MAX_cm','LENGTH_MAX_cm','YUZEY_MAX_cm2','HACIM_MAX_cm3']].head(5).to_string(index=False))

# ─── 5. Excel'e yaz ──────────────────────────────────────────
print(f"\n[5/5] Excel kaydediliyor → {OUT}")

# Renk paleti
DARK='1F3864'; MID='2E75B6'; LIGHT='D6E4F7'; WHT='FFFFFF'
ALT='EBF3FB'; GRY='F2F2F2'

t   = Side(style='thin',   color='AAAAAA')
m   = Side(style='medium', color=MID)
bdr = Border(left=t, right=t, top=t, bottom=t)
mbdr= Border(left=m, right=m, top=m, bottom=m)

def F(c, s=10, b=False):
    return Font(name='Arial', size=s, bold=b, color=c)
def P(c):
    return PatternFill('solid', fgColor=c)
def A(h='center', w=True):
    return Alignment(horizontal=h, vertical='center', wrap_text=w)

def yaz_baslik(ws, txt, n, r=1):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    c = ws.cell(row=r, column=1, value=txt)
    c.font = F(WHT, 13, True); c.fill = P(DARK)
    c.alignment = A('center', False)
    ws.row_dimensions[r].height = 26

def yaz_subhd(ws, txt, n, r):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    c = ws.cell(row=r, column=1, value=txt)
    c.font = F(WHT, 11, True); c.fill = P(MID)
    c.alignment = A('left', False)
    ws.row_dimensions[r].height = 20

def yaz_info(ws, txt, n, r, bg=LIGHT, tc='000000'):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
    c = ws.cell(row=r, column=1, value=txt)
    c.font = Font(name='Arial', size=10, color=tc, italic=True)
    c.fill = P(bg); c.alignment = A('left', True)
    c.border = mbdr; ws.row_dimensions[r].height = 30

def yaz_tablo(ws, headers, widths, data_rows, start_row):
    # Başlık satırı
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=start_row, column=i, value=h)
        c.font = F(WHT, 10, True); c.fill = P(MID)
        c.alignment = A('center', True); c.border = bdr
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[start_row].height = 32
    # Veri satırları
    for ri, row in enumerate(data_rows, start_row + 1):
        bg = ALT if ri % 2 == 0 else WHT
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = F('000000', 10); c.fill = P(bg)
            c.alignment = A('left', True); c.border = bdr
            ws.row_dimensions[ri].height = 15
            if isinstance(val, float):
                c.number_format = '#,##0.00'
    return ri  # son satır numarası

wb = openpyxl.Workbook()
wb.remove(wb.active)

# Sütun tanımları
DET_H = ["SIRA", "MALZEMENO", "PROJE", "BİRİM (orijinal→cm)",
         "GAGE (cm)", "WIDTH (cm)", "LENGTH (cm)",
         "YÜZEY ALANI (cm²)", "HACİM (cm³)", "DIMENSIONCODE"]
DET_W = [5, 30, 9, 18, 10, 10, 12, 18, 16, 22]
DET_C = ['SIRA', 'MALZEMENO', 'PROJE', 'BIRIM',
         'GAGE_cm', 'WIDTH_cm', 'LENGTH_cm',
         'YUZEY_cm2', 'HACIM_cm3', 'DIMENSIONCODE']

OZ_H  = ["MALZEMENO", "PROJE", "BİRİM",
          "MAX GAGE (cm)", "MAX WIDTH (cm)", "MAX LENGTH (cm)",
          "MAX YÜZEY (cm²)", "MAX HACİM (cm³)", "SATIR"]
OZ_W  = [30, 14, 10, 13, 13, 14, 17, 17, 7]
OZ_C  = ['MALZEMENO', 'PROJE', 'BIRIM',
          'GAGE_MAX_cm', 'WIDTH_MAX_cm', 'LENGTH_MAX_cm',
          'YUZEY_MAX_cm2', 'HACIM_MAX_cm3', 'SATIR_N']

INFO_TEXT = (
    "YÜZEY ALANI = WIDTH × LENGTH (cm²)  |  HACİM = GAGE × WIDTH × LENGTH (cm³)  |  "
    "Birim: inç×2.54=cm (GMH,LGM)  |  mm÷10=cm (WIDTH veya LENGTH >500 olan satırlar)  |  "
    "Test numuneleri (BONDTEST vb.) filtrelendi  |  SIRA 1=en büyük yüzey, SIRA 2=ikinci en büyük"
)

# ── Sekme 1: 2025 ──
NC = 10
ws1 = wb.create_sheet("2025 — Top2 Yüzey Alanı")
ws1.sheet_view.showGridLines = False
ws1.freeze_panes = 'A5'
yaz_baslik(ws1, "2025 — Her MALZEMENO için En Yüksek 2 Yüzey Alanı", NC, 1)
yaz_info(ws1, INFO_TEXT, NC, 2)
yaz_subhd(ws1,
    f"  {res25['MALZEMENO'].nunique()} unique MALZEMENO  |  {len(res25)} satır (max 2 satır/MALZEMENO)",
    NC, 3)
yaz_tablo(ws1, DET_H, DET_W, [tuple(r) for r in res25[DET_C].itertuples(index=False)], 4)

# ── Sekme 2: 2026 ──
ws2 = wb.create_sheet("2026 — Top2 Yüzey Alanı")
ws2.sheet_view.showGridLines = False
ws2.freeze_panes = 'A5'
yaz_baslik(ws2, "2026 İlk 5 Ay — Her MALZEMENO için En Yüksek 2 Yüzey Alanı", NC, 1)
yaz_info(ws2, INFO_TEXT, NC, 2)
yaz_subhd(ws2,
    f"  {res26['MALZEMENO'].nunique()} unique MALZEMENO  |  {len(res26)} satır",
    NC, 3)
yaz_tablo(ws2, DET_H, DET_W, [tuple(r) for r in res26[DET_C].itertuples(index=False)], 4)

# ── Sekme 3: Özet ──
NC3 = 9
ws3 = wb.create_sheet("ÖZET — Max Değerler")
ws3.sheet_view.showGridLines = False
ws3.freeze_panes = 'A5'
yaz_baslik(ws3, "ÖZET — MALZEMENO Başına Maksimum Yüzey Alanı & Hacim", NC3, 1)
yaz_info(ws3,
    "Her MALZEMENO için tüm satırlar içinden maksimum değerler  |  "
    "Birim dönüşümleri uygulandıktan sonra hesaplanmıştır  |  Yüzey alanına göre azalan sıra",
    NC3, 2)

yaz_subhd(ws3, f"▶ 2025 — {len(oz25)} unique MALZEMENO", NC3, 3)
last = yaz_tablo(ws3, OZ_H, OZ_W, [tuple(r) for r in oz25[OZ_C].itertuples(index=False)], 4)

yaz_subhd(ws3, f"▶ 2026 — {len(oz26)} unique MALZEMENO", NC3, last + 2)
yaz_tablo(ws3, OZ_H, OZ_W, [tuple(r) for r in oz26[OZ_C].itertuples(index=False)], last + 3)

wb.save(OUT)

print("\n✓ Tamamlandı!")
print(f"  Çıktı dosyası: {OUT}")
print(f"  Sekmeler: '2025 — Top2 Yüzey Alanı'  |  '2026 — Top2 Yüzey Alanı'  |  'ÖZET — Max Değerler'")
